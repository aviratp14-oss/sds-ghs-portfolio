# Project I: builds excel/Project_I_Data_Cleaning.xlsx, the workbook where the raw data is cleaned.
# Grey headers = raw data exactly as collected (numbers kept as text, like the source pages).
# Yellow headers = my cleaning formulas. Only plain functions (SUMIFS, VLOOKUP, SUBSTITUTE, IF), so it
# opens the same in Excel, Google Sheets and LibreOffice. 01_clean_raw_data.py is the Python cross-check.
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parents[1]
RAW = BASE / 'data' / 'raw'
OUT = BASE / 'excel' / 'Project_I_Data_Cleaning.xlsx'
OUT.parent.mkdir(exist_ok=True)

GREY, YELLOW, GREEN = (PatternFill('solid', fgColor=c) for c in ('D9D9D9', 'FFF2CC', 'E2EFDA'))
HEAD = Font(bold=True)
wb = Workbook()


def rows(path):
    return list(csv.DictReader(open(path)))


def widths(ws, w=None):
    for j in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(j)].width = (w or {}).get(j, max(11, min(45, len(str(ws.cell(1, j).value)) + 3)))


def sheet(title, data, formulas):
    """data = raw rows (grey, written as text); formulas = list of (header, template with {r})."""
    ws = wb.create_sheet(title)
    cols = list(data[0].keys())
    for j, c in enumerate(cols + [h for h, _ in formulas], 1):
        cell = ws.cell(1, j, c)
        cell.font = HEAD
        cell.fill = GREY if j <= len(cols) else YELLOW
    for i, row in enumerate(data, 2):
        for j, c in enumerate(cols, 1):
            ws.cell(i, j, int(row[c]) if c == 'Year' else row[c])
        for k, (_, f) in enumerate(formulas):
            ws.cell(i, len(cols) + 1 + k, '=' + f.format(r=i))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
    widths(ws)
    return ws, len(data) + 1


# ---------------- Notes
notes = [
    ('Project I - data cleaning: lithium demand in India', True),
    ('', False),
    ('Grey headers = raw data as collected. Yellow headers = my cleaning formulas. Green = inputs I set.', False),
    ('Sources for every row are on the Sources tab (Source_ID). A00 = my own assumption, explained next to it.', False),
    ('', False),
    ('What I fixed', True),
    ('Trade (WITS / UN Comtrade, HS 283691 lithium carbonate and 282520 lithium hydroxide, 2018-2024):', False),
    ('- numbers came as text with thousands separators and values in USD thousand: converted', False),
    ('- the World total row sits in the same list as the partners: kept it only to reconcile', False),
    ('- 16 rows have a unit value under 3 USD/kg, which no lithium chemical sells for. Ireland shows up every year', False),
    ('  at ~0.55 USD/kg; most likely a different product booked under 283691. Quantity re-estimated from value.', False),
    ('- result: reported carbonate tonnes are about 2x the cleaned figure in most years', False),
    ('- hydroxide converted at 0.880 t LCE per t LiOH.H2O so everything is in t LCE', False),
    ('EV sales (6 reports, Vahan based):', False),
    ('- 36 different segment labels mapped to 2W / 3W_L3 / 3W_L5 / 4W / Bus / Truck / CV / Total (Lookups)', False),
    ('- Indian digit grouping (20,37,831), "lakh", "~", ">" and "almost" parsed to numbers', False),
    ('- period labels (FY2025, FY 2024-25, FY2025-26, CY2025) standardised', False),
    ('- sources disagree on FY2025 (1,964,831 vs 2,037,831); kept both and noted which one I use', False),
    ('Base year 2025 units = IESA CY2025 total x segment shares; market size = EV units / EV penetration', False),
    ('Lithium per kWh worked out from cathode chemistry (formula mass, capacity, voltage) on the Chemistry tab', False),
    ('', False),
    ('The demand model itself is in Project_I_Lithium_Demand_Model.xlsx.', False),
]
ws = wb.active
ws.title = 'Notes'
for i, (t, b) in enumerate(notes, 1):
    ws.cell(i, 1, t).font = Font(bold=b, size=12 if i == 1 else 11)
ws.column_dimensions['A'].width = 115

# ---------------- Lookups
lk = wb.create_sheet('Lookups')
ev_raw = rows(RAW / 'ev_sales_reported.csv')
ev_clean = rows(BASE / 'data' / 'clean' / 'Fact_EV_Sales_Reported.csv')
seg_map = {}
for r, c in zip(ev_raw, ev_clean):
    seg_map[r['Segment_As_Reported'].strip().lower()] = c['Std_Segment']
unit_map = {}
for r, c in zip(ev_raw, ev_clean):
    unit_map[r['Unit_As_Reported'].strip().lower()] = c['Measure']
lk['A1'], lk['B1'] = 'Segment_As_Reported (lower case)', 'Std_Segment'
lk['D1'], lk['E1'] = 'Unit_As_Reported (lower case)', 'Measure'
lk['G1'], lk['H1'] = 'HS6', 't_LCE_per_t_product'
lk['J1'], lk['K1'] = 'Constant', 'Value'
for c in ('A1', 'B1', 'D1', 'E1', 'G1', 'H1', 'J1', 'K1'):
    lk[c].font = HEAD
for i, (k, v) in enumerate(sorted(seg_map.items()), 2):
    lk.cell(i, 1, k), lk.cell(i, 2, v)
SEG_N = len(seg_map) + 1
for i, (k, v) in enumerate(sorted(unit_map.items()), 2):
    lk.cell(i, 4, k), lk.cell(i, 5, v)
UNIT_N = len(unit_map) + 1
lk['G2'], lk['H2'] = '283691', 1.0
lk['G3'], lk['H3'] = '282520', 0.880
consts = [('Unit value floor (USD/kg)', 3.0), ('t LCE per t Li', 5.323), ('Electrolyte kg LCE per kWh', 0.012),
          ('Cell yield', 0.92), ('Li molar mass', 6.941), ('Fe', 55.845), ('P', 30.974), ('O', 15.999),
          ('Ni', 58.693), ('Mn', 54.938), ('Co', 58.933)]
for i, (k, v) in enumerate(consts, 2):
    lk.cell(i, 10, k)
    lk.cell(i, 11, v).fill = GREEN
C = {k: f'Lookups!$K${i}' for i, (k, _) in enumerate(consts, 2)}
widths(lk, {1: 52, 4: 30, 10: 30})

# ---------------- Raw_Trade
trade = rows(RAW / 'wits_india_imports_lithium_chemicals_2018_2024.csv')
N = len(trade) + 1
FLOOR = C['Unit value floor (USD/kg)']
tf = [
    ('Value_USD', 'VALUE(SUBSTITUTE(F{r},",",""))*1000'),
    ('Qty_kg', 'VALUE(SUBSTITUTE(G{r},",",""))'),
    ('Is_World_Row', 'E{r}="World"'),
    ('Unit_Value_USD_kg', 'IF(J{r}=0,0,I{r}/J{r})'),
    ('Flag', f'IF(K{{r}},"",IF(E{{r}}="India","Re-import",IF(L{{r}}<{FLOOR},"Unit value below floor","")))'),
    ('Use_For_Ref', 'IF(AND(NOT(K{r}),M{r}=""),1,0)'),
    ('Ref_Unit_Value', f'SUMIFS($I$2:$I${N},$B$2:$B${N},B{{r}},$C$2:$C${N},C{{r}},$N$2:$N${N},1)/SUMIFS($J$2:$J${N},$B$2:$B${N},B{{r}},$C$2:$C${N},C{{r}},$N$2:$N${N},1)'),
    ('Qty_Clean_kg', 'IF(K{r},"",IF(M{r}="Unit value below floor",I{r}/O{r},J{r}))'),
    ('Qty_Clean_t_LCE', 'IF(K{r},"",P{r}/1000*VLOOKUP(C{r},Lookups!$G$2:$H$3,2,FALSE))'),
]
sheet('Raw_Trade', trade, tf)

# ---------------- Trade_Summary
ts = wb.create_sheet('Trade_Summary')
hdr = ['Year', 'HS6', 'World_Value_USD', 'Partner_Value_USD', 'World_Qty_kg', 'Partner_Qty_kg', 'Check_Qty_Diff_kg',
       'Qty_Clean_kg', 'Qty_Removed_kg', 'Qty_Clean_t_LCE', 'Clean_Unit_Value_USD_kg']
for j, h in enumerate(hdr, 1):
    ts.cell(1, j, h).font = HEAD
    ts.cell(1, j).fill = YELLOW if j > 2 else GREY
R = "Raw_Trade!"
i = 2
for hs in ('283691', '282520'):
    for y in range(2018, 2025):
        ts.cell(i, 1, y), ts.cell(i, 2, hs)
        crit = f'{R}$B$2:$B${N},A{i},{R}$C$2:$C${N},B{i}'
        ts.cell(i, 3, f'=SUMIFS({R}$I$2:$I${N},{crit},{R}$K$2:$K${N},TRUE)')
        ts.cell(i, 4, f'=SUMIFS({R}$I$2:$I${N},{crit},{R}$K$2:$K${N},FALSE)')
        ts.cell(i, 5, f'=SUMIFS({R}$J$2:$J${N},{crit},{R}$K$2:$K${N},TRUE)')
        ts.cell(i, 6, f'=SUMIFS({R}$J$2:$J${N},{crit},{R}$K$2:$K${N},FALSE)')
        ts.cell(i, 7, f'=E{i}-F{i}')
        ts.cell(i, 8, f'=SUMIFS({R}$P$2:$P${N},{crit})')
        ts.cell(i, 9, f'=F{i}-H{i}')
        ts.cell(i, 10, f'=SUMIFS({R}$Q$2:$Q${N},{crit})')
        ts.cell(i, 11, f'=D{i}/H{i}')
        i += 1
TS_N = i - 1
ts.cell(i + 1, 1, 'Industrial baseline (avg 2022-24, t LCE)').font = HEAD
ts.cell(i + 2, 1, 'Carbonate (283691)'), ts.cell(i + 2, 2, f'=AVERAGEIFS(J2:J{TS_N},B2:B{TS_N},"283691",A2:A{TS_N},">=2022")')
ts.cell(i + 3, 1, 'Hydroxide (282520)'), ts.cell(i + 3, 2, f'=AVERAGEIFS(J2:J{TS_N},B2:B{TS_N},"282520",A2:A{TS_N},">=2022")')
CARB, HYD = f'Trade_Summary!$B${i + 2}', f'Trade_Summary!$B${i + 3}'
ts.freeze_panes = 'A2'
widths(ts, {1: 40})

# ---------------- EV_Reported
EN = len(ev_raw) + 1
ef = [
    ('Period', 'IF(LEFT(B{r},2)="CY",B{r},IF(ISNUMBER(SEARCH("target",B{r})),"Target2030",'
               'IF(LEN(SUBSTITUTE(UPPER(B{r})," ",""))=9,"FY"&MID(SUBSTITUTE(B{r}," ",""),3,2)&RIGHT(B{r},2),SUBSTITUTE(UPPER(B{r})," ",""))))'),
    ('Std_Segment', f'VLOOKUP(LOWER(TRIM(C{{r}})),Lookups!$A$2:$B${SEG_N},2,FALSE)'),
    ('Measure', f'VLOOKUP(LOWER(TRIM(E{{r}})),Lookups!$D$2:$E${UNIT_N},2,FALSE)'),
    ('Value', 'VALUE(TRIM(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(LOWER(D{r}),",",""),"~",""),">",""),"almost",""),"lakh","")))'
              '*IF(ISNUMBER(SEARCH("lakh",D{r})),100000,1)*IF(LOWER(E{r})="million",1000000,1)'),
]
sheet('EV_Reported', ev_raw, ef)

# ---------------- Base_Year_2025
by = wb.create_sheet('Base_Year_2025')
E = 'EV_Reported!'
by['A1'] = 'CY2025 total EV sales (IESA, S10)'
by['B1'] = f'=SUMIFS({E}$J$2:$J${EN},{E}$G$2:$G${EN},"CY2025",{E}$H$2:$H${EN},"Total",{E}$I$2:$I${EN},"Units")'
by['A2'] = 'Share of 3W EVs that are L3 (A00; FY2025 was 77%)'
by['B2'] = 0.75
by['B2'].fill = GREEN
heads = ['Sub_Sector', 'IESA_Segment', 'Share_of_EV_Sales', 'EV_Units_2025', 'EV_Penetration_2025 (input)', 'Market_Units_2025', 'Note']
for j, h in enumerate(heads, 1):
    by.cell(4, j, h).font = HEAD
    by.cell(4, j).fill = YELLOW
spec = [('EV_2W', '2W', 0.078, '2W penetration 6.2% in FY2025 (S08), rolled forward'),
        ('EV_3W_L5', '3W', 0.25, 'L5 penetration ~22% in FY2025 (S08)'),
        ('EV_3W_L3', '3W', 1.0, 'L3 e-rickshaws and e-carts are all electric'),
        ('EV_4W', '4W', 0.045, 'car penetration 2.7% in FY2025 and +86% sales in FY2026 (S08, S09)'),
        ('EV_Bus', 'Bus', 0.052, 'A00'),
        ('EV_Truck', 'Truck', 0.011, 'A00')]
for i, (sub, seg, pen, note) in enumerate(spec, 5):
    by.cell(i, 1, sub), by.cell(i, 2, seg)
    by.cell(i, 3, f'=SUMIFS({E}$J$2:$J${EN},{E}$G$2:$G${EN},"CY2025",{E}$H$2:$H${EN},B{i},{E}$I$2:$I${EN},"Share_of_EV_Sales")/100')
    split = '*(1-$B$2)' if sub == 'EV_3W_L5' else '*$B$2' if sub == 'EV_3W_L3' else ''
    by.cell(i, 4, f'=ROUND($B$1*C{i}{split},0)')
    by.cell(i, 5, pen).fill = GREEN
    by.cell(i, 6, f'=ROUND(D{i}/E{i},-3)')
    by.cell(i, 7, note)
by['A12'] = 'Calibration: battery GWh in 2025 at the model pack sizes vs IESA 19 GWh'
by['A12'].font = HEAD
cal = [('EV_2W', 2.8, 1), ('EV_3W_L5', 8, 1), ('EV_3W_L3', 5, 0.25), ('EV_4W', 38, 1), ('EV_Bus', 280, 1), ('EV_Truck', 40, 1)]
for j, h in enumerate(['Sub_Sector', 'kWh_per_Unit', 'LiIon_Share', 'GWh'], 1):
    by.cell(13, j, h).font = HEAD
for i, (sub, k, li) in enumerate(cal, 14):
    by.cell(i, 1, sub)
    by.cell(i, 2, k).fill = GREEN
    by.cell(i, 3, li).fill = GREEN
    by.cell(i, 4, f'=D{i - 9}*B{i}*C{i}/1000000')
by['A20'], by['D20'] = 'Model total', '=SUM(D14:D19)'
by['A21'], by['D21'] = 'IESA reported (S10)', f'=SUMIFS({E}$J$2:$J${EN},{E}$G$2:$G${EN},"CY2025",{E}$H$2:$H${EN},"Battery",{E}$I$2:$I${EN},"GWh",{E}$A$2:$A${EN},"S10")'
by['A22'], by['D22'] = 'Model as share of IESA (rest is lead-acid L3 and pack-size spread)', '=D20/D21'
widths(by, {1: 62, 7: 60})

# ---------------- Chemistry
ch = wb.create_sheet('Chemistry')
heads = ['Chemistry', 'Formula_Mass_g_mol', 'Practical_Capacity_mAh_g', 'Nominal_Voltage_V', 'Wh_per_g_CAM',
         'kg_CAM_per_kWh', 'Li_Mass_Fraction', 'kg_Li_per_kWh_Cathode', 'kg_LCE_per_kWh_Cell']
for j, h in enumerate(heads, 1):
    ch.cell(1, j, h).font = HEAD
    ch.cell(1, j).fill = YELLOW if j > 1 else GREY
L = C
chems = [('LFP  LiFePO4', f"={L['Li molar mass']}+{L['Fe']}+{L['P']}+4*{L['O']}", 160, 3.2),
         ('NMC811  LiNi0.8Mn0.1Co0.1O2', f"={L['Li molar mass']}+0.8*{L['Ni']}+0.1*{L['Mn']}+0.1*{L['Co']}+2*{L['O']}", 195, 3.7),
         ('LCO  LiCoO2', f"={L['Li molar mass']}+{L['Co']}+2*{L['O']}", 165, 3.85)]
for i, (n, mass, cap, v) in enumerate(chems, 2):
    ch.cell(i, 1, n), ch.cell(i, 2, mass)
    ch.cell(i, 3, cap).fill = GREEN
    ch.cell(i, 4, v).fill = GREEN
    ch.cell(i, 5, f'=C{i}*D{i}/1000')
    ch.cell(i, 6, f'=1/E{i}')
    ch.cell(i, 7, f"={L['Li molar mass']}/B{i}")
    ch.cell(i, 8, f'=F{i}*G{i}')
    ch.cell(i, 9, f"=(H{i}*{L['t LCE per t Li']}+{L['Electrolyte kg LCE per kWh']})/{L['Cell yield']}")
ch.cell(5, 1, 'Sodium-ion'), ch.cell(5, 9, 0)
ch['A7'] = ('Lithium per kWh = kg of cathode per kWh x lithium mass fraction, converted to LCE, plus LiPF6 salt in the '
            'electrolyte, divided by cell yield (scrap). Check: IEA and BNEF quote roughly 0.5-0.7 kg LCE per kWh.')
widths(ch, {1: 34})

# ---------------- Industry_Baseline
ib = wb.create_sheet('Industry_Baseline')
for j, h in enumerate(['Sub_Sector', 'Basis', 'Share_Input', 't_LCE_2025'], 1):
    ib.cell(1, j, h).font = HEAD
ib_rows = [('IND_Grease', 'Hydroxide imports x share to grease', 0.9, f'={HYD}*C2'),
           ('IND_GlassCeramics', 'Carbonate imports x share to glass and ceramics', 0.5, f'={CARB}*C3'),
           ('IND_Pharma_Other', 'Rest of carbonate and hydroxide', '', f'={CARB}*(1-C3)+{HYD}*(1-C2)')]
for i, (a, b, c, d) in enumerate(ib_rows, 2):
    ib.cell(i, 1, a), ib.cell(i, 2, b)
    if c != '':
        ib.cell(i, 3, c).fill = GREEN
    ib.cell(i, 4, d)
ib['A6'] = 'Cross-check: grease survey (S26)'
ib['A6'].font = HEAD
ib['A7'], ib['B7'] = 'India grease output 2020 (t) = 190m lb x 0.4536', '=190000000*0.4536/1000'
ib['A8'], ib['B8'] = 'Lithium-thickened share', 0.85
ib['A9'], ib['B9'] = 'LiOH.H2O in lithium soap grease (typical)', 0.016
ib['A10'], ib['B10'] = 'Implied t LCE (survey covers 16 of 31 makers)', '=B7*B8*B9*Lookups!$H$3'
ib['A11'], ib['B11'] = 'Model grease baseline (t LCE)', '=D2'
widths(ib, {1: 50, 2: 45})

# ---------------- Sources
src = rows(RAW / 'source_register.csv')
sheet('Sources', src, [])

# ---------------- DQ_Log
dq = wb.create_sheet('DQ_Log')
for j, h in enumerate(['Table', 'Issue', 'Rows', 'What I did'], 1):
    dq.cell(1, j, h).font = HEAD
log = [
    ('Raw_Trade', 'Numbers stored as text with separators', f'=COUNTA(Raw_Trade!F2:F{N})', 'VALUE(SUBSTITUTE()), USD thousand to USD'),
    ('Raw_Trade', 'World total rows mixed in with partners', f'=COUNTIF(Raw_Trade!K2:K{N},TRUE)', 'Used only to reconcile (Trade_Summary col G = 0)'),
    ('Raw_Trade', 'Unit value under the floor', f'=COUNTIF(Raw_Trade!M2:M{N},"Unit value below floor")', 'Quantity re-estimated from value at the clean unit value'),
    ('Raw_Trade', 'India as its own partner (re-import)', f'=COUNTIF(Raw_Trade!M2:M{N},"Re-import")', 'Kept, flagged'),
    ('Trade_Summary', 'Tonnes removed by cleaning', f'=SUM(Trade_Summary!I2:I{TS_N})/1000', 'In tonnes; most of it is carbonate'),
    ('EV_Reported', 'Segment labels needing a mapping', f'=COUNTA(Lookups!A2:A{SEG_N})', 'Lookups tab'),
    ('EV_Reported', 'Values with Indian grouping or words', f'=SUMPRODUCT(--ISNUMBER(SEARCH("lakh",EV_Reported!D2:D{EN})))+SUMPRODUCT(--ISNUMBER(SEARCH("~",EV_Reported!D2:D{EN})))+SUMPRODUCT(--ISNUMBER(SEARCH(">",EV_Reported!D2:D{EN})))+SUMPRODUCT(--ISNUMBER(SEARCH("almost",EV_Reported!D2:D{EN})))+SUMPRODUCT(--(LEN(EV_Reported!D2:D{EN})-LEN(SUBSTITUTE(EV_Reported!D2:D{EN},",",""))=2)*(MID(EV_Reported!D2:D{EN},3,1)=","))', 'Parsed to numbers'),
    ('EV_Reported', 'FY2025 total EVs: two sources disagree', 2, 'Vahan retail (S07) for history; IESA CY2025 (S10) for the base year'),
]
for i, row in enumerate(log, 2):
    for j, v in enumerate(row, 1):
        dq.cell(i, j, v)
widths(dq, {1: 18, 2: 45, 4: 70})

wb.move_sheet('DQ_Log', offset=-(len(wb.sheetnames) - 2))
wb.save(OUT)
print('saved', OUT)
