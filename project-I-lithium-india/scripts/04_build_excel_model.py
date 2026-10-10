# Project I: builds excel/Project_I_Lithium_Demand_Model.xlsx, the scenario model with toggles.
# Pick Low / Base / High on the Control tab, switch policy levers on or off, change India's accessible
# share of world supply, and every sheet recalculates. Plain formulas only (INDEX, MATCH, IF, MAX, MIN),
# so it works in Excel, Google Sheets and LibreOffice. 02_demand_supply_model.py runs the same maths in Python.
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = Path(__file__).resolve().parents[1]
CLEAN = BASE / 'data' / 'clean'
OUT = BASE / 'excel' / 'Project_I_Lithium_Demand_Model.xlsx'

import importlib.util
spec = importlib.util.spec_from_file_location('m', BASE / 'scripts' / '02_demand_supply_model.py')
import io, contextlib
with contextlib.redirect_stdout(io.StringIO()):
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)   # reuses the same assumption table, history and constants

GREY, YELLOW, GREEN, BLUE = (PatternFill('solid', fgColor=c) for c in ('D9D9D9', 'FFF2CC', 'E2EFDA', 'DDEBF7'))
HEAD, BIG = Font(bold=True), Font(bold=True, size=14)
wb = Workbook()

Y0, Y1 = 2013, 2040
def col(y):
    return get_column_letter(4 + y - Y0)          # Calc sheet: D = 2013 ... AE = 2040
def icol(y):
    return get_column_letter(8 + y - 2025)        # Inputs sheet: H = 2025 ... W = 2040

# ---------------- Control
ct = wb.active
ct.title = 'Control'
ct['A1'] = 'Lithium demand in India, 2025-2040: scenario model'
ct['A1'].font = BIG
ct['A3'], ct['B3'] = 'Adoption scenario', 'Base'
ct['A5'], ct['B5'] = 'Lever: recycling push', False
ct['A6'], ct['B6'] = 'Lever: overseas assets', False
ct['A7'], ct['B7'] = 'Lever: sodium-ion push', False
ct['A8'], ct['B8'] = 'Lever: exploration fast-track', False
ct['A10'], ct['B10'] = "India's accessible share of world lithium supply", 0.03
ct['B10'].number_format = '0.0%'
for c in ('B3', 'B5', 'B6', 'B7', 'B8', 'B10'):
    ct[c].fill = GREEN
    ct[c].font = HEAD
dv = DataValidation(type='list', formula1='"Low,Base,High"', allow_blank=False)
dv2 = DataValidation(type='list', formula1='"TRUE,FALSE"', allow_blank=False)
ct.add_data_validation(dv)
ct.add_data_validation(dv2)
dv.add('B3')
for c in ('B5', 'B6', 'B7', 'B8'):
    dv2.add(c)
notes = {3: 'Low = slow EVs and storage; Base = current trends; High = NITI Aayog-style ambition',
         5: m.LEVER_NOTE['Recycling push'], 6: m.LEVER_NOTE['Overseas assets'],
         7: m.LEVER_NOTE['Sodium-ion push'], 8: m.LEVER_NOTE['Exploration fast-track'],
         10: 'Stress test. India is ~1% of world lithium use today and ~3.5% of world GDP'}
for r, t in notes.items():
    ct.cell(r, 3, t).font = Font(italic=True, color='595959')
SCEN, LV_REC, LV_EQ, LV_NA, LV_FT, SHARE = ('Control!$B$3', 'Control!$B$5', 'Control!$B$6', 'Control!$B$7',
                                             'Control!$B$8', 'Control!$B$10')
ct.column_dimensions['A'].width = 48
ct.column_dimensions['B'].width = 12
ct.column_dimensions['C'].width = 95

# ---------------- Assumptions
asm = wb.create_sheet('Assumptions')
A = list(csv.DictReader(open(CLEAN / 'Assumptions.csv')))
heads = ['Key', 'Param', 'Sub_Sector', 'Scenario', 'Y2025', 'Y2030', 'Y2035', 'Y2040', 'Unit', 'Source_ID', 'Note']
for j, h in enumerate(heads, 1):
    asm.cell(1, j, h).font = HEAD
    asm.cell(1, j).fill = GREY if j < 5 else GREEN
for i, r in enumerate(A, 2):
    asm.cell(i, 1, f'=B{i}&"|"&C{i}&"|"&D{i}')
    asm.cell(i, 2, r['Param']), asm.cell(i, 3, r['Sub_Sector']), asm.cell(i, 4, r['Scenario'])
    for j, k in enumerate(['Y2025', 'Y2030', 'Y2035', 'Y2040'], 5):
        asm.cell(i, j, float(r[k])).fill = GREEN
    asm.cell(i, 9, r['Unit']), asm.cell(i, 10, r['Source_ID']), asm.cell(i, 11, r['Note'])
AN = len(A) + 1
asm.freeze_panes = 'E2'
for j, w in enumerate([40, 24, 18, 9, 12, 12, 12, 12, 22, 20, 90], 1):
    asm.column_dimensions[get_column_letter(j)].width = w

# ---------------- Inputs (selected scenario, interpolated by year)
inp = wb.create_sheet('Inputs')
pairs = []
for r in A:
    k = (r['Param'], r['Sub_Sector'])
    if k not in pairs:
        pairs.append(k)
for j, h in enumerate(['Param', 'Sub_Sector', 'Row_in_Assumptions', 'A2025', 'A2030', 'A2035', 'A2040'], 1):
    inp.cell(1, j, h).font = HEAD
for y in range(2025, 2041):
    inp[f'{icol(y)}1'] = y
    inp[f'{icol(y)}1'].font = HEAD
IN = {}
for i, (p, s) in enumerate(pairs, 2):
    IN[(p, s)] = i
    inp.cell(i, 1, p), inp.cell(i, 2, s)
    inp.cell(i, 3, f'=IFERROR(MATCH(A{i}&"|"&B{i}&"|"&{SCEN},Assumptions!$A$2:$A${AN},0),MATCH(A{i}&"|"&B{i}&"|All",Assumptions!$A$2:$A${AN},0))')
    for j, c in enumerate('EFGH', 4):
        inp.cell(i, j, f'=INDEX(Assumptions!${c}$2:${c}${AN},$C{i})')
    for y in range(2025, 2041):
        yc = f'{icol(y)}$1'
        inp[f'{icol(y)}{i}'] = (f'=IF({yc}<=2030,$D{i}+($E{i}-$D{i})*({yc}-2025)/5,'
                                f'IF({yc}<=2035,$E{i}+($F{i}-$E{i})*({yc}-2030)/5,$F{i}+($G{i}-$F{i})*({yc}-2035)/5))')
inp.freeze_panes = 'H2'
inp.column_dimensions['A'].width = 26
inp.column_dimensions['B'].width = 20


def I(p, s, y):
    return f'Inputs!${icol(y)}${IN[(p, s)]}'


# ---------------- History
hs = wb.create_sheet('History')
hs['A1'] = 'Sales before 2025, used only to work out what retires into recycling (A00, rounded from S07/S08 FY totals)'
hs['A1'].font = HEAD
for j, y in enumerate(m.HIST_YEARS, 2):
    hs.cell(2, j, y).font = HEAD
HR = {}
r = 3
for s in m.EV:
    hs.cell(r, 1, f'{s} units')
    for j, v in enumerate(m.HIST_UNITS[s], 2):
        hs.cell(r, j, v).fill = GREEN
    HR[s] = r
    r += 1
hs.cell(r, 1, 'EV_3W_L3 Li-ion share')
for j, v in enumerate(m.HIST_L3_LI, 2):
    hs.cell(r, j, v).fill = GREEN
HR['L3LI'] = r
r += 1
hs.cell(r, 1, 'BESS added (GWh)')
for j, v in enumerate(m.HIST_BESS_ADD, 2):
    hs.cell(r, j, v).fill = GREEN
HR['BESS'] = r
hs.column_dimensions['A'].width = 26


def H(key, y):
    return f'History!${get_column_letter(2 + y - 2013)}${HR[key]}'


# ---------------- Calc
cs = wb.create_sheet('Calc')
cs['A1'] = 'Calculation sheet: one block per use, years across. History columns (2013-2024) feed recycling only.'
cs['A1'].font = HEAD
for j, h in enumerate(['Use', 'Sub_Sector', 'Line'], 1):
    cs.cell(2, j, h).font = HEAD
for y in range(Y0, Y1 + 1):
    cs[f'{col(y)}2'] = y
    cs[f'{col(y)}2'].font = HEAD
    if y < 2025:
        cs[f'{col(y)}2'].fill = GREY
LFP_K, NMC_K, LCO_K = (m.chem[k] for k in ('LFP', 'NMC811', 'LCO'))
cs['AG2'], cs['AH2'] = 'kg LCE/kWh', ''
for i, (n, v) in enumerate([('LFP', LFP_K), ('NMC811', NMC_K), ('LCO', LCO_K)], 3):
    cs[f'AG{i}'], cs[f'AH{i}'] = n, v
    cs[f'AH{i}'].fill = GREEN
K = {'LFP': '$AH$3', 'NMC811': '$AH$4', 'LCO': '$AH$5'}
cs['AG7'] = 'from Project_I_Data_Cleaning.xlsx, Chemistry tab'

row = 3
ROW = {}


def line(use, sub, name, fn, fmt='#,##0'):
    """fn(y) -> formula string (without '=') or None."""
    global row
    cs.cell(row, 1, use), cs.cell(row, 2, sub), cs.cell(row, 3, name)
    for y in range(Y0, Y1 + 1):
        f = fn(y)
        if f is not None:
            c = cs[f'{col(y)}{row}']
            c.value = f'={f}' if isinstance(f, str) else f
            c.number_format = fmt
    ROW[(sub, name)] = row
    row += 1
    return row - 1


def R(sub, name, y):
    return f'{col(y)}{ROW[(sub, name)]}'


line('Lever', 'NaIon', 'Extra sodium-ion share (lever)',
     lambda y: None if y < 2025 else f'IF({LV_NA},IF({y}<=2030,0.15*({y}-2025)/5,IF({y}<=2035,0.15+0.15*({y}-2030)/5,0.3+0.1*({y}-2035)/5)),0)', '0.00')
row += 1
for s in m.EV:
    lab = m.LABEL[s]
    def units(y, s=s):
        return H(s, y) if y < 2025 else f'{I("Market_Units", s, y)}*{I("EV_Share", s, y)}'
    line(lab, s, 'Units', units)
    def gwh(y, s=s):
        if y < 2025:
            li = H('L3LI', y) if s == 'EV_3W_L3' else '1'
            return f'{R(s, "Units", y)}*{li}*{I("kWh_per_Unit", s, 2025)}*0.9/1000000'
        return f'{R(s, "Units", y)}*{I("LiIon_Share", s, y)}*{I("kWh_per_Unit", s, y)}/1000000'
    line(lab, s, 'GWh', gwh, '#,##0.00')
    def kk(y, s=s):
        if y < 2025:
            return f'{R(s, "kg LCE per kWh", 2025)}'
        na = f'MIN(0.8,{I("NaIon_Share", s, y)}+{R("NaIon", "Extra sodium-ion share (lever)", y)})' \
            if s in m.NA_ION_ELIGIBLE else '0'
        lfp = I('LFP_Share', s, y)
        return f'(1-{na})*({lfp}*{K["LFP"]}+(1-{lfp})*{K["NMC811"]})'
    ROW[(s, 'kg LCE per kWh')] = row       # forward ref for history columns
    line(lab, s, 'kg LCE per kWh', kk, '0.000')
    line(lab, s, 't LCE', lambda y, s=s: f'{R(s, "GWh", y)}*{R(s, "kg LCE per kWh", y)}*1000')
    row += 1

lab = m.LABEL['BESS']
line(lab, 'BESS', 'Cumulative GWh', lambda y: None if y < 2025 else I('BESS_Cumulative_GWh', 'BESS', y), '#,##0.0')
hist_sum = f'SUM(History!$B${HR["BESS"]}:$M${HR["BESS"]})'
def bess_new(y):
    # new storage ramps in a straight line inside each five-year block, so the cumulative total still hits the anchor years
    if y < 2025:
        return None
    if y == 2025:
        return f'{R("BESS", "Cumulative GWh", 2025)}-{hist_sum}'
    y0 = 2025 + 5 * ((y - 2026) // 5)
    a0, c0, c1 = f'{col(y0)}{row}', R('BESS', 'Cumulative GWh', y0), R('BESS', 'Cumulative GWh', y0 + 5)
    return f'MAX(0,{a0}+({c1}-{c0}-5*{a0})/15*({y}-{y0}))'
line(lab, 'BESS', 'New GWh (ramp)', bess_new, '#,##0.00')
def bess_add(y):
    if y < 2025:
        return H('BESS', y)
    return f'{R("BESS", "New GWh (ramp)", y)}+{col(y - 12)}{row}'
line(lab, 'BESS', 'GWh added', bess_add, '#,##0.00')
line(lab, 'BESS', 'kg LCE per kWh', lambda y: K['LFP'] if y < 2025 else
     f'(1-MIN(0.8,{I("NaIon_Share", "BESS", y)}+{R("NaIon", "Extra sodium-ion share (lever)", y)}))*{K["LFP"]}', '0.000')
line(lab, 'BESS', 't LCE', lambda y: f'{R("BESS", "GWh added", y)}*{R("BESS", "kg LCE per kWh", y)}*1000')
row += 1
for s in m.ELEC:
    lab = m.LABEL[s]
    line(lab, s, 'Units', lambda y, s=s: I('Units', s, max(y, 2025)))
    line(lab, s, 'GWh', lambda y, s=s: f'{R(s, "Units", y)}*{I("Wh_per_Unit", s, max(y, 2025))}/1000000000', '#,##0.00')
    line(lab, s, 't LCE', lambda y, s=s: f'{R(s, "GWh", y)}*{K[m.ELEC_CHEM[s]]}*1000')
row += 1
for s in ['DEF_Defence'] + m.INDS:
    line(m.LABEL[s], s, 't LCE', lambda y, s=s: None if y < 2025 else I('t_LCE', s, y))
row += 1

def total(subs):
    return lambda y: '+'.join(R(s, 't LCE', y) for s in subs) if y >= 2025 else None
line('Totals', 'Energy', 'Energy transition demand (t LCE)', total(m.EV + ['BESS']))
line('Totals', 'NonEnergy', 'Non-energy demand (t LCE)', total(m.ELEC + ['DEF_Defence'] + m.INDS))
line('Totals', 'All', 'Total demand (t LCE)', lambda y: None if y < 2025 else f'{R("Energy", "Energy transition demand (t LCE)", y)}+{R("NonEnergy", "Non-energy demand (t LCE)", y)}')
line('Totals', 'Battery', 'Battery demand incl. electronics (t LCE)',
     lambda y: None if y < 2025 else f'{R("Energy", "Energy transition demand (t LCE)", y)}+' + '+'.join(R(s, 't LCE', y) for s in m.ELEC))
row += 1

# supply
def lever_path(target, a2025):
    return lambda y: f'IF({y}<=2030,{a2025}+({target}-{a2025})*({y}-2025)/5,{target})'
def coll(param, sub, target):
    def f(y):
        if y < 2025:
            return None
        base = I(param, sub, y)
        lp = lever_path(target, I(param, sub, 2025))(y)
        return f'IF({LV_REC},MAX({base},{lp}),{base})'
    return f
line('Supply', 'Recycling', 'Collection rate EV and storage', coll('Collection_Rate', 'EV_BESS', 0.9), '0%')
line('Supply', 'Recycling', 'Collection rate electronics', coll('Collection_Rate', 'Electronics', 0.7), '0%')
line('Supply', 'Recycling', 'Lithium recovery', coll('Li_Recovery', 'Recycling', 0.95), '0%')
def recyc(y):
    if y < 2025:
        return None
    parts = []
    for s, life in m.LIFE.items():
        c = R('Recycling', 'Collection rate electronics', y) if s.startswith('ELEC') else R('Recycling', 'Collection rate EV and storage', y)
        parts.append(f'{R(s, "t LCE", y - life)}*{c}')
    return f'({"+".join(parts)})*{R("Recycling", "Lithium recovery", y)}'
line('Supply', 'Recycling', 'Recycled lithium (t LCE)', recyc)
line('Supply', 'Mining', 'Domestic mining (t LCE)', lambda y: None if y < 2025 else
     f'MAX(IF({y}>={I("Domestic_Mining_Start", "Supply", y)},{I("Domestic_Mining_t_LCE", "Supply", y)},0),'
     f'IF({LV_FT},IF({y}<2032,0,IF({y}<=2035,1500*({y}-2031)/4,1500+2500*({y}-2035)/5)),0))')
line('Supply', 'Equity', 'Overseas equity (t LCE)', lambda y: None if y < 2025 else
     f'IF({y}>={I("Overseas_Equity_Start", "Supply", y)},{I("Overseas_Equity_t_LCE", "Supply", y)},0)'
     f'+IF(AND({LV_EQ},{y}>=2031),IF({y}<=2035,10000+20000*({y}-2030)/5,30000+20000*({y}-2035)/5),0)')
line('Supply', 'Ceiling', 'Import ceiling (t LCE)', lambda y: None if y < 2025 else f'{I("Global_Supply_t_LCE", "Supply", y)}*{SHARE}')
row += 1
T = lambda n, y: R('All', 'Total demand (t LCE)', y) if n == 'tot' else None
def g(sub, name):
    return lambda y: R(sub, name, y)
line('Balance', 'Bal', 'Net import need (t LCE)', lambda y: None if y < 2025 else
     f'MAX(0,{R("All", "Total demand (t LCE)", y)}-{R("Recycling", "Recycled lithium (t LCE)", y)}-{R("Mining", "Domestic mining (t LCE)", y)})')
line('Balance', 'Bal', 'Open-market import need (t LCE)', lambda y: None if y < 2025 else
     f'MAX(0,{R("Bal", "Net import need (t LCE)", y)}-{R("Equity", "Overseas equity (t LCE)", y)})')
line('Balance', 'Bal', 'Available supply (t LCE)', lambda y: None if y < 2025 else
     f'{R("Recycling", "Recycled lithium (t LCE)", y)}+{R("Mining", "Domestic mining (t LCE)", y)}+{R("Equity", "Overseas equity (t LCE)", y)}+{R("Ceiling", "Import ceiling (t LCE)", y)}')
line('Balance', 'Bal', 'Shortfall (t LCE)', lambda y: None if y < 2025 else
     f'MAX(0,{R("All", "Total demand (t LCE)", y)}-{R("Bal", "Available supply (t LCE)", y)})')
line('Balance', 'Bal', 'Energy coverage ratio', lambda y: None if y < 2025 else
     f'MIN(1,MAX(0,{R("Bal", "Available supply (t LCE)", y)}-{R("NonEnergy", "Non-energy demand (t LCE)", y)})/{R("Energy", "Energy transition demand (t LCE)", y)})', '0%')
line('Balance', 'Bal', 'Import dependence', lambda y: None if y < 2025 else
     f'{R("Bal", "Net import need (t LCE)", y)}/{R("All", "Total demand (t LCE)", y)}', '0%')
line('Balance', 'Bal', 'Recycled content share', lambda y: None if y < 2025 else
     f'{R("Recycling", "Recycled lithium (t LCE)", y)}/{R("Battery", "Battery demand incl. electronics (t LCE)", y)}', '0.0%')
line('Balance', 'Bal', 'Recycled content mandate (BWMR)', lambda y: None if y < 2025 else m.mandate(y), '0%')
cs.freeze_panes = 'D3'
cs.column_dimensions['A'].width = 34
cs.column_dimensions['B'].width = 18
cs.column_dimensions['C'].width = 38
for y in range(Y0, Y1 + 1):
    cs.column_dimensions[col(y)].width = 10
    cs.column_dimensions[col(y)].outlineLevel = 1 if y < 2025 else 0

# ---------------- Results
rs = wb.create_sheet('Results', 1)
rs['A1'] = '=" Results: "&Control!B3&" scenario"&IF(OR(Control!B5,Control!B6,Control!B7,Control!B8),", levers on","")'
rs['A1'].font = BIG
cols = [('Year', None), ('EVs', None), ('Grid storage', None), ('Electronics', None), ('Defence and aerospace', None),
        ('Industry', None), ('Total demand', ('All', 'Total demand (t LCE)')),
        ('Energy transition', ('Energy', 'Energy transition demand (t LCE)')),
        ('Non-energy', ('NonEnergy', 'Non-energy demand (t LCE)')),
        ('Recycling', ('Recycling', 'Recycled lithium (t LCE)')), ('Domestic mining', ('Mining', 'Domestic mining (t LCE)')),
        ('Overseas equity', ('Equity', 'Overseas equity (t LCE)')), ('Import ceiling', ('Ceiling', 'Import ceiling (t LCE)')),
        ('Available supply', ('Bal', 'Available supply (t LCE)')), ('Shortfall', ('Bal', 'Shortfall (t LCE)')),
        ('Energy coverage', ('Bal', 'Energy coverage ratio')), ('Import dependence', ('Bal', 'Import dependence')),
        ('Recycled content', ('Bal', 'Recycled content share')), ('Mandate', ('Bal', 'Recycled content mandate (BWMR)'))]
groups = {'EVs': m.EV, 'Grid storage': ['BESS'], 'Electronics': m.ELEC, 'Defence and aerospace': ['DEF_Defence'], 'Industry': m.INDS}
rs['A2'] = 'All quantities in tonnes of lithium carbonate equivalent (t LCE)'
for j, (h, _) in enumerate(cols, 1):
    c = rs.cell(3, j, h)
    c.font = HEAD
    c.fill = BLUE
    c.alignment = Alignment(wrap_text=True)
for i, y in enumerate(range(2025, 2041), 4):
    rs.cell(i, 1, y)
    for j, (h, ref) in enumerate(cols[1:], 2):
        if h in groups:
            f = '=' + '+'.join(f'Calc!{R(s, "t LCE", y)}' for s in groups[h])
        else:
            f = f'=Calc!{R(ref[0], ref[1], y)}'
        c = rs.cell(i, j, f)
        c.number_format = '0%' if h in ('Energy coverage', 'Import dependence', 'Mandate') else '0.0%' if h == 'Recycled content' else '#,##0'
for j in range(1, len(cols) + 1):
    rs.column_dimensions[get_column_letter(j)].width = 12
rs.row_dimensions[3].height = 32
rs.freeze_panes = 'B4'

ch1 = BarChart()
ch1.type, ch1.grouping, ch1.overlap = 'col', 'stacked', 100
ch1.title = 'Lithium demand by use (t LCE)'
ch1.add_data(Reference(rs, min_col=2, max_col=6, min_row=3, max_row=19), titles_from_data=True)
ch1.set_categories(Reference(rs, min_col=1, min_row=4, max_row=19))
ch1.height, ch1.width = 9, 18
rs.add_chart(ch1, 'B22')
ch2 = LineChart()
ch2.title = 'Demand vs supply India can reach (t LCE)'
for c in (7, 14):
    ch2.add_data(Reference(rs, min_col=c, min_row=3, max_row=19), titles_from_data=True)
ch2.set_categories(Reference(rs, min_col=1, min_row=4, max_row=19))
ch2.height, ch2.width = 9, 18
rs.add_chart(ch2, 'L22')

wb.save(OUT)
print('saved', OUT, 'calc rows', row)
